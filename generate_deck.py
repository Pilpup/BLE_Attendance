import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="BLE_Attendance_System.pptx"):
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Professional Minimalist Tech
    BG_DARK = RGBColor(15, 23, 42)      # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)  # Slate 50
    CARD_BG = RGBColor(255, 255, 255)   # White
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    TEXT_MAIN = RGBColor(15, 23, 42)    # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    PRIMARY = RGBColor(37, 99, 235)     # Royal Blue
    SUCCESS = RGBColor(22, 163, 74)     # Emerald Green
    ALERT = RGBColor(220, 38, 38)       # Crimson Red
    PURPLE = RGBColor(124, 58, 237)     # Violet Purple
    CARD_DARK = RGBColor(30, 41, 59)    # Slate 800

    def add_header(slide, title_text, category_text="BLE ATTENDANCE SYSTEM"):
        # Category / eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme Hero)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "NEXT-GENERATION CLASSROOM VERIFICATION"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(96, 165, 250)
    p_badge.space_after = Pt(12)

    p_main = tf.add_paragraph()
    p_main.text = "Autonomous BLE Near-Field\nAttendance Verification System"
    p_main.font.size = Pt(40)
    p_main.font.bold = True
    p_main.font.color.rgb = RGBColor(255, 255, 255)
    p_main.space_after = Pt(16)

    p_sub = tf.add_paragraph()
    p_sub.text = "Zero External Hardware • Zero Manual Roll Calls • Tamper-Proof Social Mesh Graph"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_after = Pt(28)

    p_foot = tf.add_paragraph()
    p_foot.text = "Android (Kotlin + Jetpack Compose)  |  Python (FastAPI + AsyncPG)  |  Serverless PostgreSQL"
    p_foot.font.size = Pt(13)
    p_foot.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: The Problem Statement & Why Legacy Solutions Fail
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Problem: Traditional Attendance is Broken", "INDUSTRY CHALLENGE")

    col_w = Inches(3.64)
    gap = Inches(0.4)
    top_y = Inches(1.8)
    card_h = Inches(5.0)

    # Problem 1: Manual Roll Calls & Paper
    create_card(s2, Inches(0.8), top_y, col_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0), top_y + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "01. Manual Roll Calls & Paper"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ALERT
    p.space_after = Pt(14)
    bullets = [
        "Wastes 10–15 minutes of lecture time in every class.",
        "Rampant 'proxy attendance' (calling out roll for an absent friend).",
        "Sign-in paper sheets get passed around, lost, or forged.",
        "Manual data entry into college portals is tedious and error-prone."
    ]
    for b in bullets:
        p2 = tf.add_paragraph()
        p2.text = "• " + b
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(8)

    # Problem 2: QR Codes & GPS Geofencing
    create_card(s2, Inches(0.8) + col_w + gap, top_y, col_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0) + col_w + gap, top_y + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "02. QR Codes & GPS Fencing"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ALERT
    p.space_after = Pt(14)
    bullets = [
        "QR codes are photographed and instantly shared over WhatsApp.",
        "Students at home or in dorms scan the photo and get marked present.",
        "GPS does not work indoors (basement lecture halls get zero satellites).",
        "GPS margin of error (+/- 15m) cannot distinguish inside vs. outside corridor."
    ]
    for b in bullets:
        p2 = tf.add_paragraph()
        p2.text = "• " + b
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(8)

    # Problem 3: Biometrics & External Beacons
    create_card(s2, Inches(0.8) + (col_w + gap) * 2, top_y, col_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, top_y + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "03. Biometrics & Hardware Beacons"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ALERT
    p.space_after = Pt(14)
    bullets = [
        "Fingerprint scanners form massive bottlenecks (queues of 60+ students).",
        "Facial recognition is slow, privacy-invasive, and fooled by high-res photos.",
        "ESP32 / Raspberry Pi wall beacons require expensive installation & battery maintenance.",
        "Hardware beacons can be stolen, unplugged, or physically cloned."
    ]
    for b in bullets:
        p2 = tf.add_paragraph()
        p2.text = "• " + b
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 3: The Evolutionary History (What We Tried & Why It Failed)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Evolutionary Engineering: Lessons from Physical Radio Failures", "R&D PATHWAY")

    cards = [
        ("Attempt 1: 2-Beacon Differential", "Failed: The Doorway Trap", 
         "• Concept: Ratio between Teacher at podium and 1 child beacon at back wall.\n• Physical Failure: Signal ratio (+10 dB) produced identical hyperbolic curves 1m inside and 1m outside the door.\n• Conclusion: RSSI distance is scalar (no direction vector); cannot distinguish doorway spoofers."),
        
        ("Attempt 2: 3-Beacon Triangulation", "Failed: Metal Resonant Cavity", 
         "• Concept: Elect 2 corner students to triangulate classroom bounds.\n• Physical Failure: Concrete/metal reinforcement walls acted as RF waveguides. Back row received -70 dBm (reflections), middle got -82 dBm (body water absorption).\n• Conclusion: Middle student was elected as farthest beacon every time!"),
        
        ("Attempt 3: Pure BLE Multi-Hop Relay", "Failed: Packet Clashes & Buffer Limits", 
         "• Concept: Offline students broadcast 20-byte relay vouchers over air.\n• Physical Failure: Radio collision storm when 10+ students broadcast vouchers at once; 24-byte BLE payload limit; high packet loss.\n• Conclusion: Discard active multi-hop relays; adopt pure cloud-audited in-degree graph witnessing.")
    ]

    card_w3 = Inches(3.64)
    for idx, (title, subtitle, body) in enumerate(cards):
        cx = Inches(0.8) + (card_w3 + Inches(0.4)) * idx
        create_card(s3, cx, Inches(1.8), card_w3, Inches(5.0))
        tb = s3.shapes.add_textbox(cx + Inches(0.2), Inches(2.0), card_w3 - Inches(0.4), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.space_after = Pt(4)

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ALERT
        p_sub.space_after = Pt(12)

        for line in body.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(6)

    # =========================================================================
    # SLIDE 4: Our Final Solution: The Proximity vs. Transport Paradigm
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Our Final Solution: Proximity vs. Transport Separation", "CORE ARCHITECTURE")

    # Big concept card
    create_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "JOB 1: PROXIMITY SENSING (BLE)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(12)
    bullets1 = [
        "Role: Pure local physical presence radar.",
        "Answers: 'Am I physically sitting 1 meter away from Student B?'",
        "Protocol: Low-power BLE whispers (< 2.5m range, median >= -85 dBm).",
        "Duration: Exactly 5 seconds at class start.",
        "Hardware: Standard smartphone BLE radio (Zero external hardware)."
    ]
    for b in bullets1:
        p2 = tf.add_paragraph()
        p2.text = "• " + b
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(8)

    create_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tb2 = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "JOB 2: TRANSPORT & AUDIT (CLOUD API)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = SUCCESS
    p.space_after = Pt(12)
    bullets2 = [
        "Role: Secure data aggregation & mathematical proof audit.",
        "Answers: 'Does the reported topology form a legitimate cluster?'",
        "Protocol: Encrypted HTTPS REST endpoints via FastAPI.",
        "Speed: Audits 100+ students in under 15ms in cloud memory.",
        "Persistence: Permanent PostgreSQL records with roll number links."
    ]
    for b in bullets2:
        p2 = tf2.add_paragraph()
        p2.text = "• " + b
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Social Mesh Graph Engine & Mathematical Scoring
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Social Mesh Graph Engine: Bilateral Affinity Scoring", "ALGORITHM")

    # 4 Formula cards
    step_w = Inches(2.7)
    step_gap = Inches(0.28)
    for idx, (title, score, desc, color) in enumerate([
        ("Direct Mutual", "1.00 Score", "A ∈ P(B) and B ∈ P(A)\n\nBoth students reciprocally heard each other's 5s whisper. Proven physical desk neighbors.", SUCCESS),
        ("Triadic Closure", "0.50 Score", "Shares ≥ 2 verified bridges\n\nStudent heard by 2 classmates who are already verified mutuals. Propagates presence across rows.", PRIMARY),
        ("Peer Witness", "0.85 Score", "Witnessed by ≥ 2 verified peers\n\nOffline students who could not send HTTP packets, but were heard by ≥ 2 verified classmates.", PURPLE),
        ("Hallway / Spoofer", "0.00 Score", "Unreciprocated link or 0 peers\n\nAttacker claiming a student who never heard them, or isolated corridor spoofers. Flagged & rejected.", ALERT)
    ]):
        cx = Inches(0.8) + (step_w + step_gap) * idx
        create_card(s5, cx, Inches(1.8), step_w, Inches(5.0))
        tb = s5.shapes.add_textbox(cx + Inches(0.15), Inches(2.0), step_w - Inches(0.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        p_score = tf.add_paragraph()
        p_score.text = score
        p_score.font.size = Pt(20)
        p_score.font.bold = True
        p_score.font.color.rgb = color
        p_score.space_after = Pt(12)

        for line in desc.split("\n"):
            p_desc = tf.add_paragraph()
            p_desc.text = line
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = TEXT_MUTED if not line else TEXT_MAIN
            p_desc.space_after = Pt(4)

    # =========================================================================
    # SLIDE 6: 3-Layer Security & Anti-Spoofing Architecture
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "3-Layer Security Architecture: Defense-in-Depth", "SECURITY MODEL")

    layers = [
        ("Layer 1: Cryptographic Digital Identity", "JWT BEARER TOKENS", 
         "• Proves: 'Who is holding this phone account?'\n• Cryptographically signed with HS256 secret.\n• Enforces strict role authorization (Teacher vs Student).\n• Prevents submitting attendance under another student's ID."),
        
        ("Layer 2: In-Room Time-Synchronized Nonce", "DYNAMIC BLE NONCE (SALT)", 
         "• Proves: 'Is this phone physically inside the classroom right now?'\n• Teacher broadcasts ephemeral rolling salt over BLE (S:<SESSION>:<SALT>).\n• Cannot be guessed by students at home.\n• Backend checks salt; mismatches rejected as 'Nice try lil bro'."),
        
        ("Layer 3: Physical Desk Proximity Mesh", "LOW-POWER BLE WHISPERS", 
         "• Proves: 'Are real classmates physically surrounding this phone?'\n• Transmitted strictly at TX_POWER_LOW (-15 dBm, <= 2.5m fence).\n• Walls cause >20 dB attenuation, eliminating corridor spoofers.\n• Requires in-degree mutual witnessing in graph topology.")
    ]

    for idx, (title, tag, body) in enumerate(layers):
        cx = Inches(0.8) + (step_w + step_gap) * idx
        # widen to 3 columns
        col3_w = Inches(3.64)
        cx3 = Inches(0.8) + (col3_w + Inches(0.4)) * idx
        create_card(s6, cx3, Inches(1.8), col3_w, Inches(5.0))
        tb = s6.shapes.add_textbox(cx3 + Inches(0.2), Inches(2.0), col3_w - Inches(0.4), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN
        p_t.space_after = Pt(12)

        for line in body.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(6)

    # =========================================================================
    # SLIDE 7: Offline Capabilities: Zero-Packet Pure Graph Witnessing
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Offline Capabilities: Autonomous Presence in Dead Zones", "OFFLINE SAFETY NET")

    create_card(s7, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tb = s7.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "How a Student with Zero Internet / Data Plan is Verified Present"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    steps = [
        ("1. In-Room Radio Reception", "The offline student's phone receives the teacher's BLE beacon (S:89F2:9381) via direct radio waves (zero internet required)."),
        ("2. Near-Field Whisper Probe", "The offline phone broadcasts its 8-byte whisper (89F2OFF1) at TX_POWER_LOW for 5 seconds."),
        ("3. Verified Online Neighbors Catch Probe", "Classmates to their left and right (who have cellular data) hear the probe at >= -85 dBm and add 'OFF1' to their Top-3 neighbors."),
        ("4. Pure In-Degree Cloud Audit", "The online neighbors submit their lists. The cloud graph engine observes: 'OFF1 submitted no HTTP report, but is witnessed by 2 verified classmates (ST01 and ST02)'. OFF1 is marked VERIFIED_PRESENT (0.85 affinity)."),
        ("5. Single-Witness Safety Net", "If an offline student sits alone in a corner with only 1 online witness, they are marked ISOLATED_FLAGGED. The teacher confirms visually and verifies them in 1 tap.")
    ]

    for title, desc in steps:
        p_s = tf.add_paragraph()
        p_s.text = title + " — "
        p_s.font.size = Pt(13)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_MAIN

        p_s2 = tf.add_paragraph()
        p_s2.text = desc
        p_s2.font.size = Pt(12)
        p_s2.font.color.rgb = TEXT_MUTED
        p_s2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: Frontend Android Architecture (Kotlin + Jetpack Compose)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Frontend Architecture: Native Android (Kotlin & Compose)", "MOBILE CLIENT")

    col_w4 = Inches(2.7)
    col_gap4 = Inches(0.28)
    cards_fe = [
        ("Strict Role Isolation", "• Zero manual mode toggle.\n• UI locks to Teacher or Student based on authenticated JWT claims.\n• Teachers never scan; students never see teacher console."),
        ("BLE Hardware Layer", "• BleAdvertiser: Low-latency service data advertising within 24-byte ceiling.\n• BleScanner: Hardware UUID scan filter in LOW_LATENCY mode.\n• Noise rejection filters ambient Bluetooth packets in 0ms."),
        ("Jetpack Compose UI", "• Modern Material 3 reactive layout.\n• Dual-mode Teacher Console: Broadcast Setup mode & Collapsible Audit Results.\n• Student telemetry card with live dBm signal meters."),
        ("State Management", "• AttendanceViewModel manages 5-state student machine (Standby -> Probing -> Verified).\n• Auto session countdown & graceful HTTP error recovery.\n• Live Compose snapshot reactivity.")
    ]

    for idx, (title, body) in enumerate(cards_fe):
        cx = Inches(0.8) + (col_w4 + col_gap4) * idx
        create_card(s8, cx, Inches(1.8), col_w4, Inches(5.0))
        tb = s8.shapes.add_textbox(cx + Inches(0.15), Inches(2.0), col_w4 - Inches(0.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p.space_after = Pt(12)

        for line in body.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(6)

    # =========================================================================
    # SLIDE 9: Backend Architecture (FastAPI + Serverless PostgreSQL)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Backend Architecture: High-Throughput Cloud Engine", "CLOUD BACKEND")

    cards_be = [
        ("FastAPI Async Engine", "• Python 3.11 with ASGI asynchronous concurrency.\n• Sub-15ms response latency under heavy classroom load.\n• Clean modular routing (/auth, /attendance)."),
        ("PostgreSQL on Neon", "• Serverless PostgreSQL with auto-scaling compute.\n• AsyncPG driver with PgBouncer pooling.\n• 5 relational tables: Users, Courses, Sessions, Records, Staging Submissions."),
        ("Ephemeral Staging Table", "• Multi-worker crash resilience.\n• Check-ins staged in session_submissions DB table instead of volatile server RAM.\n• Process-isolated across Uvicorn worker threads."),
        ("Single-Roundtrip Migration", "• Atomic DO $$ BEGIN...END $$; block at application startup.\n• Eliminates 35+ WAN roundtrips on cold start.\n• Production-sanitized error handler prevents credential leakage.")
    ]

    for idx, (title, body) in enumerate(cards_be):
        cx = Inches(0.8) + (col_w4 + col_gap4) * idx
        create_card(s9, cx, Inches(1.8), col_w4, Inches(5.0))
        tb = s9.shapes.add_textbox(cx + Inches(0.15), Inches(2.0), col_w4 - Inches(0.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = SUCCESS
        p.space_after = Pt(12)

        for line in body.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(6)

    # =========================================================================
    # SLIDE 10: Performance, Optimizations & Test Verification
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Optimizations, Testing & Key Milestones", "VERIFICATION")

    col_w2 = Inches(5.6)
    gap2 = Inches(0.5)

    # Card 1: Key Optimizations
    create_card(s10, Inches(0.8), Inches(1.8), col_w2, Inches(5.0))
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(2.0), col_w2 - Inches(0.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Engineering Optimizations"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(12)
    opts = [
        "5-Second Whisper Window: Reduced radio transmission time from 18.7s to 5.0s, reducing smartphone battery draw by 73%.",
        "Calibrated Proximity Threshold: Tuned median RSSI cutoff to -85 dBm for accurate classroom desk separation without hallway leakage.",
        "Automatic Session Supersede: Expired / abandoned sessions are auto-cancelled to eliminate teacher zombie lockouts.",
        "Auto-Logout on Expiration: Seamless 401 handling gracefully prompts for re-login with 1-year token persistence.",
        "Compose Snapshot Mutation Fix: Optimized peer map updates so RSSI meters and packet counters live-recompose instantly."
    ]
    for o in opts:
        p2 = tf.add_paragraph()
        p2.text = "• " + o
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(6)

    # Card 2: Test Suite & Simulation
    create_card(s10, Inches(0.8) + col_w2 + gap2, Inches(1.8), col_w2, Inches(5.0))
    tb2 = s10.shapes.add_textbox(Inches(1.1) + col_w2 + gap2, Inches(2.0), col_w2 - Inches(0.5), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Rigorous Testing & E2E Simulation"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = SUCCESS
    p.space_after = Pt(12)
    tests = [
        "100% Pytest Suite Passing: 12 comprehensive unit and integration test suites covering auth, graph algorithms, and security.",
        "30-Student E2E Classroom Mesh Simulation: Simulated 25 online students, 4 offline students, and 1 hallway spoofer across 6 desk clusters.",
        "Simulation Result: Exactly 29 students verified (including all 4 offline students), 1 hallway spoofer rejected with 0.00 score.",
        "Anti-Cheat Verification: Verified that unreciprocated hallway links and invalid salt attempts are trapped and flagged on teacher screen.",
        "Android Build Success: Tested on physical Android devices and compiling cleanly under Gradle 8.7."
    ]
    for t in tests:
        p2 = tf2.add_paragraph()
        p2.text = "• " + t
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_after = Pt(6)

    # =========================================================================
    # SLIDE 11: Summary & Conclusion (Dark Theme)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    bg11 = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg11.fill.solid()
    bg11.fill.fore_color.rgb = BG_DARK
    bg11.line.fill.background()

    tbox11 = s11.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.8))
    tf11 = tbox11.text_frame
    tf11.word_wrap = True

    p_badge11 = tf11.paragraphs[0]
    p_badge11.text = "PROJECT SUMMARY & VISION"
    p_badge11.font.size = Pt(13)
    p_badge11.font.bold = True
    p_badge11.font.color.rgb = RGBColor(96, 165, 250)
    p_badge11.space_after = Pt(12)

    p_head11 = tf11.add_paragraph()
    p_head11.text = "Frictionless, Autonomous & Tamper-Proof"
    p_head11.font.size = Pt(36)
    p_head11.font.bold = True
    p_head11.font.color.rgb = RGBColor(255, 255, 255)
    p_head11.space_after = Pt(20)

    summary_points = [
        "Eliminates 100% of attendance roll call waste: Classes start in 5 seconds while phones rest in pockets.",
        "Zero hardware budget required: Runs solely on existing smartphones with standard Bluetooth Low Energy.",
        "Robust against fraud: Multi-layer security defeats buddy punchers, hallway proxies, and WhatsApp QR forwarders.",
        "Zero-configuration deployment: Works in any lecture hall without seating charts, room measurements, or calibration.",
        "Resilient offline safety net: Peer witnessing ensures students in basement halls never lose attendance credits."
    ]
    for pt in summary_points:
        p_pt = tf11.add_paragraph()
        p_pt.text = "✓  " + pt
        p_pt.font.size = Pt(15)
        p_pt.font.color.rgb = RGBColor(226, 232, 240)
        p_pt.space_after = Pt(10)

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "/home/wink/Projects/ble_attendance/BLE_Attendance_System.pptx"
    create_deck(out)
